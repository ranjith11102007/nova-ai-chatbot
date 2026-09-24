import io
import logging

from .config import settings

logger = logging.getLogger("nova_ai")

ALLOWED_DOC_EXTENSIONS = {"txt", "pdf", "docx"}


def extract_text(filename: str, content_type: str, data: bytes) -> str:
    """Extract readable text from an uploaded document.

    Raises a RuntimeError with a friendly message when the file cannot be read.
    """
    name = (filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""

    if ext not in ALLOWED_DOC_EXTENSIONS:
        raise RuntimeError("This file type isn't supported.")

    try:
        if ext == "txt":
            text = _extract_txt(data)
        elif ext == "pdf":
            text = _extract_pdf(data)
        else:
            text = _extract_docx(data)
    except RuntimeError:
        raise
    except Exception as exc:
        logger.warning("Could not extract '%s': %s", filename, exc)
        raise RuntimeError("Sorry, I couldn't process this file.") from None

    text = (text or "").strip()
    if not text:
        raise RuntimeError(
            "Sorry, I couldn't extract any text from this file."
        )

    return text[: settings.MAX_DOCUMENT_CHARS]


def _extract_txt(data: bytes) -> str:
    # Try UTF-8 first, then fall back to a permissive decode.
    try:
        return data.decode("utf-8", errors="strict")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts).strip()